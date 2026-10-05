import csv
import json
import os
import re
import subprocess
import sys
from datetime import datetime
 
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CASE_DIR = os.path.join(BASE_DIR, "case")
CSV_FILE = os.path.join(CASE_DIR, "case.csv")
XLSX_FILE = os.path.join(CASE_DIR, "case.xlsx")
EXPORT_FIELDS = ["case_id", "timestamp", "user_input", "ai_response"]
 
CASE_PATTERN = re.compile(r"case_(\d+)\.json")
 
 
# --- Case files ---
def list_case_files():
    """Returns paths of all case_<n>.json files, sorted by number (oldest first)."""
    found = []
    if os.path.isdir(CASE_DIR):
        for name in os.listdir(CASE_DIR):
            match = CASE_PATTERN.fullmatch(name)
            if match:
                found.append((int(match.group(1)), os.path.join(CASE_DIR, name)))
    return [path for _, path in sorted(found)]
 
 
def case_path(case_id):
    return os.path.join(CASE_DIR, f"{case_id}.json")
 
 
def load_case(filepath):
    """Loads one case file and returns it as a dict, or None if it can't be read."""
    if not os.path.exists(filepath):
        return None
 
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            data = json.load(f)
    except (json.JSONDecodeError, OSError) as e:
        print(f"[Warning] Could not load '{filepath}': {e}")
        return None
 
    case_id = os.path.splitext(os.path.basename(filepath))[0]
 
    if isinstance(data, list):  # legacy format: just a list of records
        return {"case_id": case_id, "status": "closed", "logs": data}
 
    if isinstance(data, dict):
        case = dict(data)
        case["case_id"] = case_id
        case.setdefault("status", "closed")
        case.setdefault("logs", [])
        return case
 
    return None
 
 
def save_case(case):
    """Writes a case to disk (via a temp file so a crash can't corrupt it)."""
    os.makedirs(CASE_DIR, exist_ok=True)
    path = case_path(case["case_id"])
    tmp_path = path + ".tmp"
 
    try:
        with open(tmp_path, "w", encoding="utf-8") as f:
            json.dump(case, f, indent=4, ensure_ascii=False)
        os.replace(tmp_path, path)
        return True
    except OSError as e:
        print(f"[Warning] Failed to save '{path}': {e}")
        return False
 
 
def create_new_case():
    """Creates the next numbered case file and returns it (status 'open')."""
    numbers = [
        int(CASE_PATTERN.fullmatch(os.path.basename(p)).group(1))
        for p in list_case_files()
    ]
    case = {
        "case_id": f"case_{max(numbers, default=0) + 1}",
        "status": "open",
        "created": datetime.now().isoformat(),
        "logs": [],
    }
    save_case(case)
    return case
 
 
def get_active_case():
    """Returns the newest case if it is still open, otherwise None."""
    files = list_case_files()
    if not files:
        return None
    case = load_case(files[-1])
    return case if case and case["status"] == "open" else None
 
 
def close_case(case):
    case["status"] = "closed"
    case["closed"] = datetime.now().isoformat()
    return save_case(case)
 
 
def add_log(case, user_input, response_text):
    """Appends one exchange to the case and saves it."""
    case["logs"].append({
        "timestamp": datetime.now().isoformat(),
        "user_input": user_input,
        "ai_response": response_text,
    })
    return save_case(case)
 
 
# --- Querying records across all cases ---
def load_all_logs():
    """Returns every log from every case, each tagged with its case_id."""
    records = []
    for path in list_case_files():
        case = load_case(path)
        if case:
            for log in case["logs"]:
                records.append({**log, "case_id": case["case_id"]})
    return records
 
 
def filter_records_by_keyword(keyword, records=None):
    """Finds records with the keyword in the user input or the AI response."""
    if records is None:
        records = load_all_logs()
 
    keyword_lower = keyword.lower()
    return [
        r for r in records
        if keyword_lower in r.get("user_input", "").lower()
        or keyword_lower in r.get("ai_response", "").lower()
    ]
 
 
# --- Exporting ---
def convert_json_to_xlsx(xlsx_file=XLSX_FILE):
    """Writes all saved records to an Excel workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Side, PatternFill
        from openpyxl.formatting.rule import FormulaRule
    except ImportError:
        print("Excel export needs the 'openpyxl' package.")
        print("Install it with: pip install openpyxl")
        print("Or use the 'export' command to save a CSV instead, which Excel can also open.")
        return None

    records = load_all_logs()
    if not records:
        print("No records to export yet.")
        return None

    os.makedirs(os.path.dirname(xlsx_file), exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Assessments"
    ws.append(EXPORT_FIELDS)
    for record in records:
        ws.append([record.get(field, "") for field in EXPORT_FIELDS])

    for column, width in zip("ABCD", (12, 25, 50, 80)):
        ws.column_dimensions[column].width = width

    # Wrap text in every cell (top-aligned so tall rows read cleanly)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(EXPORT_FIELDS)):
        for cell in row:
            cell.alignment = wrap
        # Bold line under the last row of each case (wherever case_id changes)
    thick = Border(bottom=Side(style="medium"))
    for i in range(len(records) - 1):
        if records[i].get("case_id") != records[i + 1].get("case_id"):
            for cell in ws[i + 2][:len(EXPORT_FIELDS)]:  # +2 = header row + 1-based index
                cell.border = thick

    # Conditional formatting on ai_response (shades the whole cell)
    col = chr(ord("A") + EXPORT_FIELDS.index("ai_response"))
    rng = f"{col}2:{col}{ws.max_row}"
    first = f"{col}2"
    rules = [  # (text the response starts with, fill colour)
        ("urgency classification: insufficient_information", "808080"),  # grey
        ("urgency classification: emergency",                "FFC7CE"),  # red
        ("urgency classification: non-urgent",               "FFFF00"),  # yellow
        ("urgency classification: urgent",                   "FFA500"),  # orange
    ]
    for text, colour in rules:
        fill = PatternFill(start_color=colour, end_color=colour, fill_type="solid")
        ws.conditional_formatting.add(
            rng,
            FormulaRule(formula=[f'LEFT(LOWER({first}),{len(text)})="{text}"'], fill=fill),
        )

    try:
        wb.save(xlsx_file)
    except OSError as e:
        print(f"Error: could not write '{xlsx_file}': {e}")
        return None

    print(f"Exported {len(records)} record(s) to '{xlsx_file}'.")
    return len(records)
 
 
def convert_json_to_csv(csv_file=CSV_FILE):
    """Writes all saved records (from every case) to a CSV file."""
    records = load_all_logs()
    if not records:
        print("No records to export yet.")
        return None
    os.makedirs(os.path.dirname(csv_file), exist_ok=True)
 
    try:
        # utf-8-sig lets Excel read the characters correctly;
        # newline="" stops blank lines appearing on Windows.
        with open(csv_file, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=EXPORT_FIELDS, extrasaction="ignore")
            writer.writeheader()
            for record in records:
                writer.writerow({field: record.get(field, "") for field in EXPORT_FIELDS})
    except OSError as e:
        print(f"Error: could not write '{csv_file}': {e}")
        return None
 
    print(f"Exported {len(records)} record(s) to '{csv_file}'.")
    return len(records)
 
 
def open_file(filename=XLSX_FILE):
    """Opens a file with the system's default application."""
    if not os.path.exists(filename):
        print(f"'{filename}' not found. Type 'export xlsx' first to create it.")
        return False
 
    try:
        if sys.platform.startswith("win"):
            os.startfile(filename)
        elif sys.platform == "darwin":
            subprocess.run(["open", filename], check=True)
        else:
            subprocess.run(["xdg-open", filename], check=True)
    except (OSError, subprocess.CalledProcessError) as e:
        print(f"Could not open '{filename}': {e}")
        return False
 
    print(f"Opening '{filename}'...")
    return True
 
 
# --- History browsing ---
def show_history(results, page_size=2):
    """Paginated viewer for a list of records."""
    total_records = len(results)
    total_pages = (total_records + page_size - 1) // page_size
    current_page = 0
 
    while True:
        start_idx = current_page * page_size
        page_items = results[start_idx:start_idx + page_size]
 
        print(f"\n{'=' * 20} History (Page {current_page + 1} of {total_pages}) {'=' * 20}")
 
        for idx, record in enumerate(page_items, start=start_idx + 1):
            print(f"\n[Record #{idx}/{total_records}] {record.get('case_id', '')} - {record['timestamp']}")
            print(f"User Question:\n{record['user_input']}")
            print(f"\nFull AI Response:\n{record['ai_response']}")
            print("-" * 60)
 
        nav_options = []
        if current_page < total_pages - 1:
            nav_options.append("'n' for next page")
        if current_page > 0:
            nav_options.append("'p' for previous page")
        nav_options.append("'q' to exit search")
 
        action = input(f"\nNavigation [{', '.join(nav_options)}]: ").strip().lower()
 
        if action == "n" and current_page < total_pages - 1:
            current_page += 1
        elif action == "p" and current_page > 0:
            current_page -= 1
        elif action in ("q", "exit", "quit"):
            print("Exited history view.")
            break
        else:
            print("Invalid option or end of pages reached.")