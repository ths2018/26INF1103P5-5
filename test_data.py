"""
Offline tests for data_manager.py

- No real data touched: CASE_DIR is redirected to a temporary directory for
  every test, so your real 'case/' folder is never read or modified.
- No dependency on input_output_manager being importable: a lightweight stub
  module is registered in sys.modules BEFORE data_manager is imported.
- Standard library only (unittest + unittest.mock), so it also runs under
  pytest. The XLSX tests additionally need 'openpyxl' and are skipped if it
  is not installed.

Run:
    python -m unittest -v test_data
"""

import csv
import json
import os
import shutil
import sys
import tempfile
import types
import unittest
from unittest import mock

# ---------------------------------------------------------------------------
# Make data_manager importable whether this file sits next to it or in tests/
# ---------------------------------------------------------------------------
_HERE = os.path.dirname(os.path.abspath(__file__))
for _path in (_HERE, os.path.dirname(_HERE)):
    if _path not in sys.path:
        sys.path.insert(0, _path)

# ---------------------------------------------------------------------------
# Stub the one module data_manager imports from
# ---------------------------------------------------------------------------
_stub = types.ModuleType("input_output_manager")
_stub.print_output = mock.MagicMock(name="input_output_manager.print_output")
sys.modules["input_output_manager"] = _stub

# If another test file (e.g. test_logic_manager) already registered a fake
# "data_manager", drop it so we test the real module.
sys.modules.pop("data_manager", None)

import data_manager as dm  # noqa: E402  (must come after the stub)

try:
    import openpyxl  # noqa: F401
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ---------------------------------------------------------------------------
# Hardcoded exchanges (same layout the real AI is prompted to produce)
# ---------------------------------------------------------------------------
Q_HEADACHE = "I have a headache and a mild fever"
Q_ANKLE = "My ankle is swollen"
Q_CHEST = "I have chest pain"
Q_UNICODE = "头痛 – café"

R_NON_URGENT = "Urgency classification: non-urgent\nCertainty: 85%\n\nRest and drink fluids."
R_URGENT = "Urgency classification: urgent\nCertainty: 70%\n\nSee a doctor within 24 hours."
R_EMERGENCY = "Urgency classification: emergency\nCertainty: 95%\n\nCall emergency services now."
R_MULTILINE = 'Line one, with comma\nLine "two" with quotes'


class LegacyLog:
    """One log record as stored inside a case file."""

    @staticmethod
    def make(user_input, ai_response, timestamp="2026-01-01T10:00:00"):
        return {"timestamp": timestamp, "user_input": user_input, "ai_response": ai_response}


class DataManagerTestCase(unittest.TestCase):
    """Base class: redirects CASE_DIR to a temp dir and silences output."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.case_dir = os.path.join(self.tmp, "case")
        self.csv_file = os.path.join(self.case_dir, "case.csv")
        self.xlsx_file = os.path.join(self.case_dir, "case.xlsx")

        patcher = mock.patch.object(dm, "CASE_DIR", self.case_dir)
        patcher.start()
        self.addCleanup(patcher.stop)

        out = mock.patch.object(dm.input_output_manager, "print_output", mock.MagicMock())
        self.print_output = out.start()
        self.addCleanup(out.stop)

        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    # --- helpers ---
    def write_raw(self, name, content):
        """Writes a raw file into the temp case folder and returns its path."""
        os.makedirs(self.case_dir, exist_ok=True)
        path = os.path.join(self.case_dir, name)
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return path

    def make_case(self, exchanges=None, close=False):
        """Creates a real case on disk, optionally with logs and closed."""
        case = dm.create_new_case()
        for user_input, ai_response in exchanges or []:
            dm.add_log(case, user_input, ai_response)
        if close:
            dm.close_case(case)
        return case

    def printed(self):
        """All text passed to print_output, joined together."""
        return "\n".join(
            str(c.args[0]) for c in self.print_output.call_args_list if c.args
        )


# ---------------------------------------------------------------------------
# Case files: listing and paths
# ---------------------------------------------------------------------------
class TestListCaseFiles(DataManagerTestCase):
    def test_missing_directory_returns_empty_list(self):
        self.assertEqual(dm.list_case_files(), [])

    def test_sorted_numerically_not_alphabetically(self):
        for n in (10, 2, 1):
            self.write_raw(f"case_{n}.json", "{}")
        names = [os.path.basename(p) for p in dm.list_case_files()]
        self.assertEqual(names, ["case_1.json", "case_2.json", "case_10.json"])

    def test_ignores_non_matching_files(self):
        self.write_raw("case_1.json", "{}")
        self.write_raw("case_1.json.tmp", "{}")
        self.write_raw("case.csv", "")
        self.write_raw("notes.json", "{}")
        self.write_raw("case_abc.json", "{}")
        names = [os.path.basename(p) for p in dm.list_case_files()]
        self.assertEqual(names, ["case_1.json"])

    def test_case_path_appends_json(self):
        self.assertEqual(
            dm.case_path("case_3"), os.path.join(self.case_dir, "case_3.json")
        )


# ---------------------------------------------------------------------------
# Case files: loading and saving
# ---------------------------------------------------------------------------
class TestLoadCase(DataManagerTestCase):
    def test_missing_file_returns_none(self):
        self.assertIsNone(dm.load_case(os.path.join(self.case_dir, "case_1.json")))

    def test_invalid_json_returns_none_and_warns(self):
        path = self.write_raw("case_1.json", "{not valid json")
        self.assertIsNone(dm.load_case(path))
        self.assertIn("[Warning] Could not load", self.printed())

    def test_legacy_list_format_is_wrapped_and_closed(self):
        logs = [LegacyLog.make(Q_HEADACHE, R_NON_URGENT)]
        path = self.write_raw("case_4.json", json.dumps(logs))
        case = dm.load_case(path)
        self.assertEqual(case["case_id"], "case_4")
        self.assertEqual(case["status"], "closed")
        self.assertEqual(case["logs"], logs)

    def test_dict_gets_defaults(self):
        path = self.write_raw("case_2.json", json.dumps({"created": "x"}))
        case = dm.load_case(path)
        self.assertEqual(case["status"], "closed")
        self.assertEqual(case["logs"], [])
        self.assertEqual(case["created"], "x")

    def test_case_id_comes_from_filename_not_content(self):
        path = self.write_raw(
            "case_7.json", json.dumps({"case_id": "wrong", "status": "open"})
        )
        case = dm.load_case(path)
        self.assertEqual(case["case_id"], "case_7")
        self.assertEqual(case["status"], "open")

    def test_unsupported_json_type_returns_none(self):
        path = self.write_raw("case_1.json", json.dumps("just a string"))
        self.assertIsNone(dm.load_case(path))


class TestSaveCase(DataManagerTestCase):
    def test_roundtrip(self):
        case = {"case_id": "case_1", "status": "open", "logs": []}
        self.assertTrue(dm.save_case(case))
        self.assertEqual(dm.load_case(dm.case_path("case_1")), case)

    def test_creates_directory_if_missing(self):
        self.assertFalse(os.path.isdir(self.case_dir))
        dm.save_case({"case_id": "case_1", "status": "open", "logs": []})
        self.assertTrue(os.path.isdir(self.case_dir))

    def test_no_temp_file_left_behind(self):
        dm.save_case({"case_id": "case_1", "status": "open", "logs": []})
        self.assertFalse(any(n.endswith(".tmp") for n in os.listdir(self.case_dir)))

    def test_unicode_preserved(self):
        case = {
            "case_id": "case_1", "status": "open",
            "logs": [LegacyLog.make(Q_UNICODE, R_NON_URGENT)],
        }
        dm.save_case(case)
        loaded = dm.load_case(dm.case_path("case_1"))
        self.assertEqual(loaded["logs"][0]["user_input"], Q_UNICODE)

    def test_failure_returns_false_and_warns(self):
        with mock.patch.object(dm.os, "replace", side_effect=OSError("disk full")):
            ok = dm.save_case({"case_id": "case_1", "status": "open", "logs": []})
        self.assertFalse(ok)
        self.assertIn("[Warning] Failed to save", self.printed())


# ---------------------------------------------------------------------------
# Case lifecycle: create, active, close, log
# ---------------------------------------------------------------------------
class TestCreateNewCase(DataManagerTestCase):
    def test_first_case_is_case_1_and_open(self):
        case = dm.create_new_case()
        self.assertEqual(case["case_id"], "case_1")
        self.assertEqual(case["status"], "open")
        self.assertEqual(case["logs"], [])
        self.assertIn("created", case)
        self.assertTrue(os.path.exists(dm.case_path("case_1")))

    def test_numbers_increment(self):
        dm.create_new_case()
        self.assertEqual(dm.create_new_case()["case_id"], "case_2")

    def test_uses_highest_number_even_with_gaps(self):
        self.write_raw("case_5.json", "{}")
        self.assertEqual(dm.create_new_case()["case_id"], "case_6")


class TestActiveAndClose(DataManagerTestCase):
    def test_no_cases_means_no_active_case(self):
        self.assertIsNone(dm.get_active_case())

    def test_newest_open_case_is_active(self):
        created = dm.create_new_case()
        self.assertEqual(dm.get_active_case()["case_id"], created["case_id"])

    def test_closed_newest_case_is_not_active(self):
        dm.close_case(dm.create_new_case())
        self.assertIsNone(dm.get_active_case())

    def test_only_newest_case_is_considered(self):
        old = dm.create_new_case()  # stays open
        dm.close_case(dm.create_new_case())
        self.assertIsNone(dm.get_active_case())
        self.assertEqual(dm.load_case(dm.case_path(old["case_id"]))["status"], "open")

    def test_close_case_sets_status_and_timestamp_and_persists(self):
        case = dm.create_new_case()
        self.assertTrue(dm.close_case(case))
        loaded = dm.load_case(dm.case_path(case["case_id"]))
        self.assertEqual(loaded["status"], "closed")
        self.assertIn("closed", loaded)


class TestAddLog(DataManagerTestCase):
    def test_appends_and_persists(self):
        case = dm.create_new_case()
        self.assertTrue(dm.add_log(case, Q_HEADACHE, R_NON_URGENT))
        loaded = dm.load_case(dm.case_path(case["case_id"]))
        self.assertEqual(len(loaded["logs"]), 1)
        log = loaded["logs"][0]
        self.assertEqual(log["user_input"], Q_HEADACHE)
        self.assertEqual(log["ai_response"], R_NON_URGENT)
        self.assertIn("timestamp", log)

    def test_order_is_preserved(self):
        case = dm.create_new_case()
        for q in (Q_HEADACHE, Q_ANKLE, Q_CHEST):
            dm.add_log(case, q, R_NON_URGENT)
        loaded = dm.load_case(dm.case_path(case["case_id"]))
        self.assertEqual(
            [log["user_input"] for log in loaded["logs"]],
            [Q_HEADACHE, Q_ANKLE, Q_CHEST],
        )


# ---------------------------------------------------------------------------
# Querying records across cases
# ---------------------------------------------------------------------------
class TestLoadAllLogs(DataManagerTestCase):
    def test_empty_when_no_cases(self):
        self.assertEqual(dm.load_all_logs(), [])

    def test_logs_tagged_with_case_id_across_cases(self):
        self.make_case([(Q_HEADACHE, R_NON_URGENT), (Q_ANKLE, R_URGENT)], close=True)
        self.make_case([(Q_CHEST, R_EMERGENCY)])
        records = dm.load_all_logs()
        self.assertEqual(
            [r["case_id"] for r in records], ["case_1", "case_1", "case_2"]
        )
        self.assertEqual(
            [r["user_input"] for r in records], [Q_HEADACHE, Q_ANKLE, Q_CHEST]
        )

    def test_corrupt_case_is_skipped(self):
        self.make_case([(Q_HEADACHE, R_NON_URGENT)], close=True)
        self.write_raw("case_2.json", "garbage")
        self.assertEqual(len(dm.load_all_logs()), 1)


class TestFilterRecords(DataManagerTestCase):
    def setUp(self):
        super().setUp()
        self.records = [
            {"user_input": "Chest pain", "ai_response": "Urgency classification: EMERGENCY"},
            {"user_input": "Mild cough", "ai_response": "Urgency classification: non-urgent"},
            {"user_input": "Rash", "ai_response": "See a doctor about the chest"},
        ]

    def test_matches_user_input_case_insensitively(self):
        result = dm.filter_records_by_keyword("COUGH", self.records)
        self.assertEqual([r["user_input"] for r in result], ["Mild cough"])

    def test_matches_ai_response(self):
        self.assertEqual(len(dm.filter_records_by_keyword("emergency", self.records)), 1)

    def test_matches_either_field(self):
        self.assertEqual(len(dm.filter_records_by_keyword("chest", self.records)), 2)

    def test_no_match_returns_empty(self):
        self.assertEqual(dm.filter_records_by_keyword("zzz", self.records), [])

    def test_defaults_to_all_saved_logs(self):
        self.make_case([(Q_HEADACHE, R_NON_URGENT), (Q_ANKLE, R_URGENT)])
        self.assertEqual(len(dm.filter_records_by_keyword("ankle")), 1)

    def test_missing_fields_do_not_crash(self):
        result = dm.filter_records_by_keyword("x", [{"user_input": "x"}, {}])
        self.assertEqual(len(result), 1)


# ---------------------------------------------------------------------------
# Exporting: CSV
# ---------------------------------------------------------------------------
class TestCsvExport(DataManagerTestCase):
    def test_no_records_returns_none_and_writes_nothing(self):
        self.assertIsNone(dm.convert_json_to_csv(self.csv_file))
        self.assertFalse(os.path.exists(self.csv_file))
        self.assertIn("No records to export yet.", self.printed())

    def test_exports_header_and_rows(self):
        self.make_case([(Q_HEADACHE, R_NON_URGENT), (Q_ANKLE, R_URGENT)], close=True)
        self.make_case([(Q_CHEST, R_EMERGENCY)])

        self.assertEqual(dm.convert_json_to_csv(self.csv_file), 3)

        with open(self.csv_file, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual(list(rows[0].keys()), dm.EXPORT_FIELDS)
        self.assertEqual(
            [r["user_input"] for r in rows], [Q_HEADACHE, Q_ANKLE, Q_CHEST]
        )
        self.assertEqual([r["case_id"] for r in rows], ["case_1", "case_1", "case_2"])
        self.assertIn("Exported 3 record(s)", self.printed())

    def test_file_has_utf8_bom_for_excel(self):
        self.make_case([(Q_HEADACHE, R_NON_URGENT)])
        dm.convert_json_to_csv(self.csv_file)
        with open(self.csv_file, "rb") as f:
            self.assertEqual(f.read(3), b"\xef\xbb\xbf")

    def test_multiline_comma_and_quote_content_roundtrips(self):
        self.make_case([(Q_UNICODE, R_MULTILINE)])
        dm.convert_json_to_csv(self.csv_file)
        with open(self.csv_file, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual(rows[0]["user_input"], Q_UNICODE)
        self.assertEqual(rows[0]["ai_response"], R_MULTILINE)

    def test_write_error_returns_none(self):
        # Records are supplied directly so that only the CSV write hits the
        # patched open(); otherwise reading the case JSON would fail first.
        record = {**LegacyLog.make(Q_HEADACHE, R_NON_URGENT), "case_id": "case_1"}
        with mock.patch.object(dm, "load_all_logs", return_value=[record]), \
             mock.patch("builtins.open", side_effect=OSError("denied")):
            self.assertIsNone(dm.convert_json_to_csv(self.csv_file))
        self.assertIn("Error: could not write", self.printed())


# ---------------------------------------------------------------------------
# Exporting: XLSX
# ---------------------------------------------------------------------------
class TestXlsxExport(DataManagerTestCase):
    def test_missing_openpyxl_prints_help_and_returns_none(self):
        self.make_case([(Q_HEADACHE, R_NON_URGENT)])
        real_import = __import__

        def fake_import(name, *args, **kwargs):
            if name.startswith("openpyxl"):
                raise ImportError("no openpyxl")
            return real_import(name, *args, **kwargs)

        with mock.patch("builtins.__import__", side_effect=fake_import):
            self.assertIsNone(dm.convert_json_to_xlsx(self.xlsx_file))
        self.assertIn("python -m pip install openpyxl", self.printed())

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl not installed")
    def test_no_records_returns_none(self):
        self.assertIsNone(dm.convert_json_to_xlsx(self.xlsx_file))
        self.assertFalse(os.path.exists(self.xlsx_file))

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl not installed")
    def test_workbook_contents_and_formatting(self):
        from openpyxl import load_workbook

        self.make_case(
            [(Q_CHEST, R_EMERGENCY), (Q_HEADACHE, R_NON_URGENT)], close=True
        )
        self.make_case([(Q_ANKLE, R_URGENT)])

        self.assertEqual(dm.convert_json_to_xlsx(self.xlsx_file), 3)
        ws = load_workbook(self.xlsx_file).active

        self.assertEqual(ws.title, "Assessments")
        self.assertEqual([c.value for c in ws[1]], dm.EXPORT_FIELDS)
        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws["C2"].value, Q_CHEST)
        self.assertEqual(ws["A4"].value, "case_2")

        # Text wrapping applied
        self.assertTrue(ws["D2"].alignment.wrap_text)

        # Thick border only under the last row of case_1 (row 3)
        self.assertEqual(ws["A3"].border.bottom.style, "medium")
        self.assertIsNone(ws["A2"].border.bottom.style)
        self.assertIsNone(ws["A4"].border.bottom.style)

        # One range with four urgency colour rules
        rules = [r for cf in ws.conditional_formatting for r in cf.rules]
        self.assertEqual(len(rules), 4)

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl not installed")
    def test_save_error_returns_none(self):
        self.make_case([(Q_HEADACHE, R_NON_URGENT)])
        with mock.patch("openpyxl.Workbook.save", side_effect=OSError("locked")):
            self.assertIsNone(dm.convert_json_to_xlsx(self.xlsx_file))
        self.assertIn("Error: could not write", self.printed())


# ---------------------------------------------------------------------------
# open_file
# ---------------------------------------------------------------------------
class TestOpenFile(DataManagerTestCase):
    def existing_file(self):
        path = os.path.join(self.tmp, "export.xlsx")
        open(path, "w").close()
        return path

    def test_missing_file_returns_false(self):
        self.assertFalse(dm.open_file(os.path.join(self.tmp, "nope.xlsx")))
        self.assertIn("not found", self.printed())

    def test_linux_uses_xdg_open(self):
        path = self.existing_file()
        with mock.patch.object(dm.sys, "platform", "linux"), \
             mock.patch.object(dm.subprocess, "run") as run:
            self.assertTrue(dm.open_file(path))
        run.assert_called_once_with(["xdg-open", path], check=True)

    def test_macos_uses_open(self):
        path = self.existing_file()
        with mock.patch.object(dm.sys, "platform", "darwin"), \
             mock.patch.object(dm.subprocess, "run") as run:
            self.assertTrue(dm.open_file(path))
        run.assert_called_once_with(["open", path], check=True)

    def test_windows_uses_startfile(self):
        path = self.existing_file()
        with mock.patch.object(dm.sys, "platform", "win32"), \
             mock.patch.object(dm.os, "startfile", create=True) as startfile:
            self.assertTrue(dm.open_file(path))
        startfile.assert_called_once_with(path)

    def test_launcher_failure_returns_false(self):
        path = self.existing_file()
        error = dm.subprocess.CalledProcessError(1, "xdg-open")
        with mock.patch.object(dm.sys, "platform", "linux"), \
             mock.patch.object(dm.subprocess, "run", side_effect=error):
            self.assertFalse(dm.open_file(path))
        self.assertIn("Could not open", self.printed())


# ---------------------------------------------------------------------------
# show_history (paginated viewer; page_size defaults to 2)
# ---------------------------------------------------------------------------
class TestShowHistory(DataManagerTestCase):
    def setUp(self):
        super().setUp()
        # 5 records -> 3 pages at the default page size of 2
        self.results = [
            {
                "case_id": "case_1",
                "timestamp": f"t{i}",
                "user_input": f"question {i}",
                "ai_response": f"answer {i}",
            }
            for i in range(1, 6)
        ]

    def run_history(self, *keys):
        """Runs show_history with scripted keypresses."""
        with mock.patch("builtins.input", side_effect=list(keys)):
            dm.show_history(self.results)
        return self.printed()

    def test_quit_immediately_shows_first_page_only(self):
        out = self.run_history("q")
        self.assertIn("Page 1 of 3", out)
        self.assertIn("question 1", out)
        self.assertIn("question 2", out)
        self.assertNotIn("question 3", out)
        self.assertIn("Exited history view.", out)

    def test_next_then_previous(self):
        out = self.run_history("n", "p", "q")
        self.assertIn("Page 2 of 3", out)
        self.assertIn("question 3", out)
        self.assertEqual(out.count("Page 1 of 3"), 2)

    def test_cannot_go_previous_on_first_page(self):
        self.assertIn("Invalid option", self.run_history("p", "q"))

    def test_cannot_go_past_last_page(self):
        out = self.run_history("n", "n", "n", "q")
        self.assertIn("Page 3 of 3", out)
        self.assertIn("Invalid option", out)

    def test_last_page_shows_remaining_record(self):
        self.assertIn("[Record #5/5]", self.run_history("n", "n", "q"))

    def test_quit_aliases(self):
        for word in ("exit", "quit", "Q"):
            with self.subTest(word=word):
                self.assertIn("Exited history view.", self.run_history(word))


if __name__ == "__main__":
    unittest.main(verbosity=2)